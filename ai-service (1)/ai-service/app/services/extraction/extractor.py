"""Text extraction from PDF (PyMuPDF), DOCX (python-docx), and plain text/HTML."""
from __future__ import annotations

import base64
import binascii
import io
import re
from dataclasses import dataclass, field
from html.parser import HTMLParser
from typing import List, Optional

from app.services.validation.injection import detect_injection

MAX_BYTES = 25 * 1024 * 1024


class ExtractionError(ValueError):
    pass


@dataclass
class Extracted:
    text: str
    mime_type: str
    page_count: Optional[int] = None
    warnings: List[str] = field(default_factory=list)
    injection_signals: List[str] = field(default_factory=list)


def _sniff(data: bytes, filename: str, mime: Optional[str]) -> str:
    if data[:5] == b"%PDF-":
        return "application/pdf"
    if data[:2] == b"PK" and filename.lower().endswith(".docx"):
        return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if filename.lower().endswith((".html", ".htm")) or (mime or "").endswith("html"):
        return "text/html"
    if filename.lower().endswith((".txt", ".md")) or (mime or "").startswith("text/"):
        return "text/plain"
    raise ExtractionError("unsupported or unrecognized file type")


def _pdf(data: bytes) -> tuple[str, int, List[str]]:
    import fitz  # PyMuPDF

    warnings: List[str] = []
    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception as e:  # fitz raises various types for corrupt files
        raise ExtractionError("could not open PDF") from e
    if doc.needs_pass:
        raise ExtractionError("PDF is encrypted")
    pages = [p.get_text("text") for p in doc]
    n = len(pages)
    text = "\n\n".join(pages).strip()
    if not text:
        warnings.append("No extractable text; document may be scanned (OCR not performed).")
    return text, n, warnings


def _docx(data: bytes) -> tuple[str, List[str]]:
    import docx

    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as e:
        raise ExtractionError("could not open DOCX") from e
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts).strip(), []


class _HTMLText(HTMLParser):
    SKIP = {"script", "style", "noscript"}

    def __init__(self):
        super().__init__()
        self.out: List[str] = []
        self._skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self._skip += 1
        if tag in {"p", "br", "div", "li", "h1", "h2", "h3", "h4", "tr"}:
            self.out.append("\n")

    def handle_endtag(self, tag):
        if tag in self.SKIP and self._skip:
            self._skip -= 1

    def handle_data(self, d):
        if not self._skip:
            self.out.append(d)


def extract_bytes(data: bytes, filename: str, mime: Optional[str] = None) -> Extracted:
    if len(data) > MAX_BYTES:
        raise ExtractionError("file too large")
    kind = _sniff(data, filename, mime)
    warnings: List[str] = []
    pages = None
    if kind == "application/pdf":
        text, pages, warnings = _pdf(data)
    elif kind.endswith("wordprocessingml.document"):
        text, warnings = _docx(data)
    elif kind == "text/html":
        p = _HTMLText()
        p.feed(data.decode("utf-8", errors="replace"))
        text = re.sub(r"\n{3,}", "\n\n", "".join(p.out)).strip()
    else:
        text = data.decode("utf-8", errors="replace").strip()
    signals = detect_injection(text)
    if signals:
        warnings.append("Document contains instruction-like text; it will be treated strictly as data.")
    return Extracted(text=text, mime_type=kind, page_count=pages, warnings=warnings, injection_signals=signals)


def extract_base64(content_b64: str, filename: str, mime: Optional[str] = None) -> Extracted:
    try:
        data = base64.b64decode(content_b64, validate=True)
    except (binascii.Error, ValueError) as e:
        raise ExtractionError("invalid base64 content") from e
    return extract_bytes(data, filename, mime)
